import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

class ExcelLogic:
    @staticmethod
    def get_clean_value(campo, row_data):
        val = row_data.get(campo, "")
        if not val:
            c_clean = campo.strip().lower().rstrip(":")
            for k, v in row_data.items():
                if k.strip().lower().rstrip(":") == c_clean:
                    val = v
                    break
            if not val:
                return ""
            
        campo_lower = campo.strip().lower().replace("é", "e").replace("ê", "e").rstrip(":")
        if campo_lower not in ["genero", "sexo"]:
            return str(val)

        matching_key = None
        if "_options" in row_data and isinstance(row_data["_options"], dict):
            c_clean = campo.strip().lower().replace("é", "e").replace("ê", "e").rstrip(":")
            for opt_key in row_data["_options"].keys():
                opt_key_clean = opt_key.strip().lower().replace("é", "e").replace("ê", "e").rstrip(":")
                if opt_key_clean == c_clean or (c_clean in ["genero", "sexo"] and opt_key_clean in ["genero", "sexo"]):
                    matching_key = opt_key
                    break
                    
        if matching_key:
            opt_list = row_data["_options"][matching_key]
            if isinstance(opt_list, list) and opt_list:
                indices = [x.strip() for x in str(val).split(",") if x.strip()]
                cleaned_opts = []
                for idx_str in indices:
                    matched = False
                    for opt in opt_list:
                        opt_clean = opt.strip()
                        if opt_clean.lower().startswith(idx_str.lower() + ".") or \
                           opt_clean.lower().startswith(idx_str.lower() + "-") or \
                           opt_clean.lower().startswith(idx_str.lower() + " "):
                            txt = opt_clean.split(".", 1)[-1].split("-", 1)[-1].strip()
                            if " " in txt and txt.split(" ", 1)[0].isdigit():
                                txt = txt.split(" ", 1)[-1].strip()
                            cleaned_opts.append(txt)
                            matched = True
                            break
                        elif idx_str.lower() in opt_clean.lower():
                            txt = opt_clean
                            if "." in txt: txt = txt.split(".", 1)[-1]
                            elif "-" in txt: txt = txt.split("-", 1)[-1]
                            cleaned_opts.append(txt.strip())
                            matched = True
                            break
                    if not matched:
                        cleaned_opts.append(idx_str)
                return ", ".join(cleaned_opts)
        return str(val)

    @staticmethod
    def format_header(name: str) -> str:
        words = name.split()
        if len(words) <= 1:
            return name
        abbr = words[0][:3] + "."
        return abbr + " " + " ".join(words[1:])

    @staticmethod
    def generate_excel(columns, extracted_rows, save_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Datos extraidos"

        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill("solid", fgColor="1a7a4a")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin = Side(border_style="thin", color="AAAAAA")
        cell_border = Border(left=thin, right=thin, bottom=thin, top=thin)

        headers = [ExcelLogic.format_header(c) for c in columns]
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = cell_border
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = max(len(header) + 4, 14)

        ws.row_dimensions[1].height = 22

        data_align = Alignment(horizontal="left", vertical="center")
        for row_idx, row_data in enumerate(extracted_rows, start=2):
            for col_idx, col_name in enumerate(columns, start=1):
                val = ExcelLogic.get_clean_value(col_name, row_data)
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.alignment = data_align
                cell.border = cell_border

        ws.freeze_panes = "A2"

        all_options = {}
        for row_data in extracted_rows:
            if "_options" in row_data and isinstance(row_data["_options"], dict):
                for k, v in row_data["_options"].items():
                    if k not in all_options and isinstance(v, list) and v:
                        all_options[k] = v

        if all_options:
            options_col_idx = len(columns) + 2
            current_row = 1
            
            option_header_font = Font(bold=True, size=11)
            option_text_font = Font(size=10)
            
            for param, opt_list in all_options.items():
                cell = ws.cell(row=current_row, column=options_col_idx, value=f"*{param}")
                cell.font = option_header_font
                current_row += 1
                
                for opt in opt_list:
                    cell = ws.cell(row=current_row, column=options_col_idx, value=str(opt))
                    cell.font = option_text_font
                    current_row += 1
                
                current_row += 1

            max_opt_len = 15
            for r in range(1, current_row):
                val = ws.cell(row=r, column=options_col_idx).value
                if val:
                    max_opt_len = max(max_opt_len, len(str(val)))
            ws.column_dimensions[openpyxl.utils.get_column_letter(options_col_idx)].width = max_opt_len + 4

        wb.save(save_path)
