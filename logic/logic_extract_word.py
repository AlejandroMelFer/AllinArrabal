import re
import os
import tempfile
import zipfile
from docx import Document
from docx.text.paragraph import Paragraph

class WordLogic:
    @staticmethod
    def is_checkbox_cell(text):
        text_clean = text.replace(" ", "")
        if "☐" in text or "☑" in text or "☒" in text or "□" in text: return True
        if "[ ]" in text or "[x]" in text.lower() or "[✓]" in text: return True
        if "\u00a8" in text or "\u00fe" in text or "\u00fd" in text: return True
        return False

    @staticmethod
    def is_writable(text):
        t = text.strip()
        return t == "" or t == "..." or all(c == "." or c == "_" or c == " " for c in t)

    @staticmethod
    def clean_label(text):
        t = text.replace("\r", " ").replace("\n", " ")
        while "  " in t: t = t.replace("  ", " ")
        t = t.strip()
        for sym in ["☐", "☑", "☒", "□", "[ ]", "[x]", "[X]", "[✓]", "\u00a8", "\u00fe", "\u00fd"]:
            t = t.replace(sym, "")
        return t.strip().rstrip(":").strip()

    @staticmethod
    def extract_placeholders(docx_path):
        doc = Document(docx_path)
        
        # Guardar depuración en un archivo local para poder leer la estructura de tablas directamente
        try:
            debug_path = os.path.join(os.getcwd(), "debug_table.txt")
            with open(debug_path, "w", encoding="utf-8") as df:
                df.write("--- DEBUG WORD TABLES START ---\n")
                for table_idx, table in enumerate(doc.tables):
                    df.write(f"Table {table_idx}:\n")
                    for row_idx, row in enumerate(table.rows):
                        cells_strs = []
                        for col_idx, cell in enumerate(row.cells):
                            # Limpiar saltos de línea para ver en una sola línea en la depuración
                            clean_cell_text = cell.text.replace('\n', ' ').replace('\r', '')
                            cells_strs.append(f"C{col_idx}:'{clean_cell_text}'")
                        df.write(f"  R{row_idx}: {' | '.join(cells_strs)}\n")
        except Exception as e:
            print(f"Error al escribir debug_table.txt: {e}")

        placeholders = []
        placeholders_map = {}
        regex_placeholders = set()
        
        pattern = re.compile(r'\{\{([^}]+)\}\}')
        
        namespace = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        try:
            p_elements = doc.element.body.findall('.//w:p', namespace)
            for p_elem in p_elements:
                p = Paragraph(p_elem, doc)
                for m in pattern.findall(p.text):
                    field = m.strip()
                    if field not in placeholders: placeholders.append(field)
                    regex_placeholders.add(field)
        except Exception:
            for p in doc.paragraphs:
                for m in pattern.findall(p.text):
                    field = m.strip()
                    if field not in placeholders: placeholders.append(field)
                    regex_placeholders.add(field)

        processed_coords = set()
        for table_idx, table in enumerate(doc.tables):
            row_idx = 0
            num_rows = len(table.rows)
            while row_idx < num_rows:
                row = table.rows[row_idx]
                if len(row.cells) < 2:
                    row_idx += 1; continue
                    
                label_text = row.cells[0].text.strip()
                if not label_text or WordLogic.is_writable(label_text):
                    row_idx += 1; continue
                    
                merged_rows = [row_idx]
                next_row_idx = row_idx + 1
                while next_row_idx < num_rows:
                    next_row = table.rows[next_row_idx]
                    if len(next_row.cells) >= 2 and next_row.cells[0].text.strip() == label_text:
                        merged_rows.append(next_row_idx)
                        next_row_idx += 1
                    else: break
                        
                if len(merged_rows) > 1:
                    field_name = WordLogic.clean_label(label_text)
                    if field_name:
                        if field_name not in placeholders: placeholders.append(field_name)
                        if field_name not in placeholders_map: placeholders_map[field_name] = []
                            
                        for r_idx in merged_rows:
                            val_cell_text = table.rows[r_idx].cells[1].text.strip()
                            clean_opt = val_cell_text
                            for sym in ["☐", "☑", "☒", "□", "[ ]", "[x]", "[X]", "[✓]", "\u00a8", "\u00fe", "\u00fd"]:
                                clean_opt = clean_opt.replace(sym, "")
                            placeholders_map[field_name].append((table_idx, r_idx, 1, 'merged_checkbox', clean_opt.strip()))
                            processed_coords.update([(table_idx, r_idx, 0), (table_idx, r_idx, 1)])
                    row_idx = next_row_idx
                else: row_idx += 1

        for table_idx, table in enumerate(doc.tables):
            for row_idx, row in enumerate(table.rows):
                for col_idx in range(len(row.cells) - 1):
                    if (table_idx, row_idx, col_idx) in processed_coords: continue
                        
                    label_text = row.cells[col_idx].text.strip()
                    val_text = row.cells[col_idx + 1].text.strip()
                    
                    if label_text and not WordLogic.is_writable(label_text):
                        field_name = WordLogic.clean_label(label_text)
                        if not field_name: continue
                            
                        if WordLogic.is_writable(val_text):
                            if field_name not in placeholders: placeholders.append(field_name)
                            if field_name not in placeholders_map: placeholders_map[field_name] = []
                            placeholders_map[field_name].append((table_idx, row_idx, col_idx + 1, 'text', ''))
                        elif WordLogic.is_checkbox_cell(val_text):
                            if field_name not in placeholders: placeholders.append(field_name)
                            if field_name not in placeholders_map: placeholders_map[field_name] = []
                            placeholders_map[field_name].append((table_idx, row_idx, col_idx + 1, 'checkbox', ''))
                            
        return placeholders, placeholders_map, regex_placeholders

    @staticmethod
    def get_clean_value(campo, row_data):
        val = row_data.get(campo, "")
        if not val: return ""
            
        if "_options" in row_data and isinstance(row_data["_options"], dict) and campo in row_data["_options"]:
            opt_list = row_data["_options"][campo]
            if isinstance(opt_list, list) and opt_list:
                indices = [x.strip() for x in str(val).split(",") if x.strip()]
                resolved = []
                for idx in indices:
                    prefix_dot, prefix_space = f"{idx}.", f"{idx} "
                    found = False
                    for opt in opt_list:
                        opt_str = str(opt).strip()
                        if opt_str.startswith(prefix_dot) or opt_str.startswith(prefix_space):
                            clean_opt = opt_str[len(prefix_dot):].strip() if opt_str.startswith(prefix_dot) else opt_str[len(prefix_space):].strip()
                            resolved.append(clean_opt)
                            found = True
                            break
                    if not found: resolved.append(idx)
                return ", ".join(resolved)
        return str(val)

    @staticmethod
    def check_checkboxes(cell_text, clean_val):
        has_ballot = "☐" in cell_text or "☑" in cell_text or "☒" in cell_text
        has_bracket = "[ ]" in cell_text or "[x]" in cell_text.lower() or "[✓]" in cell_text
        has_wingdings = "\u00a8" in cell_text or "\u00fe" in cell_text or "\u00fd" in cell_text

        selected = str(clean_val).lower().strip()

        if has_ballot:
            cell_text = cell_text.replace("☑", "☐").replace("☒", "☐")
            parts, empty_sym, check_sym = cell_text.split("☐"), "☐", "☒"
        elif has_bracket:
            cell_text = cell_text.replace("[x]", "[ ]").replace("[X]", "[ ]").replace("[✓]", "[ ]")
            parts, empty_sym, check_sym = cell_text.split("[ ]"), "[ ]", "[X]"
        elif has_wingdings:
            cell_text = cell_text.replace("\u00fe", "\u00a8").replace("\u00fd", "\u00a8")
            parts, empty_sym, check_sym = cell_text.split("\u00a8"), "\u00a8", "\u00fd"
        else: return cell_text

        new_parts = [parts[0]]
        for idx, part in enumerate(parts[1:]):
            opt_clean = part.strip().lower()
            if selected == opt_clean or selected == str(idx + 1):
                new_parts.append(f"{check_sym} " + part.lstrip())
            else:
                new_parts.append(f"{empty_sym} " + part.lstrip())
        return "".join(new_parts)

    @staticmethod
    def generate_word_documents(extracted_rows, template_path, placeholders_map, regex_placeholders, dest_zip):
        success_count = 0
        with tempfile.TemporaryDirectory() as temp_dir:
            for idx, row_data in enumerate(extracted_rows):
                nombre = row_data.get("Nombre", "").strip()
                apellido = row_data.get("Apellido 1", "").strip() or row_data.get("Apellidos", "").strip()
                
                out_filename = f"Ficha_{nombre}_{apellido}".strip("_") + ".docx" if nombre or apellido else f"Documento_Extraido_{idx + 1}.docx"
                out_filename = "".join(c for c in out_filename if c.isalnum() or c in "._- ")
                temp_file_path = os.path.join(temp_dir, out_filename)

                doc = Document(template_path)
                
                for field_name, coords in placeholders_map.items():
                    clean_val = WordLogic.get_clean_value(field_name, row_data)
                    for table_idx, row_idx, col_idx, field_type, option_text in coords:
                        if table_idx < len(doc.tables) and row_idx < len(doc.tables[table_idx].rows) and col_idx < len(doc.tables[table_idx].rows[row_idx].cells):
                            cell = doc.tables[table_idx].rows[row_idx].cells[col_idx]
                            if field_type == 'checkbox':
                                cell.text = WordLogic.check_checkboxes(cell.text, clean_val)
                            elif field_type == 'merged_checkbox':
                                is_selected = False
                                opt_clean, sel_clean = option_text.strip().lower(), str(clean_val).strip().lower()
                                
                                siblings = [x for x in coords if x[3] == 'merged_checkbox']
                                current_sibling_idx = next((i for i, sib in enumerate(siblings) if sib[1] == row_idx), -1)
                                
                                if sel_clean == opt_clean or sel_clean == str(current_sibling_idx + 1) or (sel_clean and (sel_clean in opt_clean or opt_clean in sel_clean)):
                                    is_selected = True
                                    
                                cell.text = f"{'☒' if is_selected else '☐'} {option_text.strip()}"
                            else:
                                cell.text = clean_val

                if regex_placeholders:
                    for p in doc.paragraphs:
                        for run in p.runs:
                            for ph in regex_placeholders:
                                tag = f"{{{{{ph}}}}}"
                                if tag in run.text:
                                    run.text = run.text.replace(tag, WordLogic.get_clean_value(ph, row_data))

                    for table in doc.tables:
                        for row in table.rows:
                            for cell in row.cells:
                                for p in cell.paragraphs:
                                    for run in p.runs:
                                        for ph in regex_placeholders:
                                            tag = f"{{{{{ph}}}}}"
                                            if tag in run.text:
                                                run.text = run.text.replace(tag, WordLogic.get_clean_value(ph, row_data))
                doc.save(temp_file_path)
                success_count += 1

            with zipfile.ZipFile(dest_zip, 'w', zipfile.ZIP_DEFLATED) as zipf:
                for root, _, files in os.walk(temp_dir):
                    for file in files:
                        zipf.write(os.path.join(root, file), arcname=file)
                        
        return success_count
